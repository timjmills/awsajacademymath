"""Pacing workbook (xlsx) for K-G5 built from the preview data. No exam question content."""
import sys, json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
sys.argv = [sys.argv[0], 'inputs/adapt.json']
import gen_docs as G   # reuses the data model (also rewrites the HTML, harmless)

F = 'Arial'; NAVY = '1F3864'
HDR = PatternFill('solid', fgColor=NAVY); HF = Font(name=F, bold=True, color='FFFFFF', size=10)
TH = Side(style='thin', color='D0D0D0'); BOX = Border(left=TH, right=TH, top=TH, bottom=TH)
DCOL = {'CC': 'FFF2CC', 'OA': 'FCE4D6', 'NBT': 'E2EFDA', 'NF': 'EDE1F5', 'MD': 'DDEBF7', 'G': 'E4DFF5', '': 'EDEDED'}
EXAMF = PatternFill('solid', fgColor='F4CCCC'); MAPF = PatternFill('solid', fgColor='D9C8EC')
KIND = {'core': 'White Rose', 'copy': 'COPIED IN', 'build': 'CCSS BUILD (to be made)', 'enrich': 'Enrichment (after MAP)'}


def hdr(ws, cols, widths):
    for j, (h, w) in enumerate(zip(cols, widths), 1):
        c = ws.cell(1, j, h); c.fill = HDR; c.font = HF; c.alignment = Alignment(wrap_text=True, vertical='center')
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.row_dimensions[1].height = 30; ws.freeze_panes = 'A2'


wb = Workbook()
ws = wb.active; ws.title = 'READ ME'; ws.column_dimensions['A'].width = 120
L = [('Awsaj Math · K to Grade 5 taught by CCSS domain · 2026-27', 16, True),
     ('Built from the curriculum site preview, so it matches the site week for week.', 10, False), ('', 10, False),
     ('Each Common Core domain is taught as one unbroken run and closed by its Awsaj domain exam (red rows), on every strand in the same week.', 10, False),
     ('Every grade-level standard is taught before the spring MAP Growth window (purple row); enrichment lessons come after it.', 10, False),
     ('Three strands share the map: Standard (~4 new lessons a week), Priority (~3) and Intervention (~2). The S / P / I columns show which strand teaches the lesson in that week.', 10, False),
     ('COPIED IN = a lesson from another grade, added so every exam question is taught first; teach it from the source grade\'s files.', 10, False),
     ('CCSS BUILD = an Awsaj lesson for a standard White Rose does not teach in this grade; it still has to be written.', 10, False),
     ('Support = the lower-grade lessons (Pre-K up to the grade below) listed for that week\'s support blocks.', 10, False),
     ('Exams sat by W14 count for the Semester 1 report; the rest for Semester 2.', 10, False)]
for i, (t, sz, b) in enumerate(L, 1):
    c = ws.cell(i, 1, t); c.font = Font(name=F, size=sz, bold=b, color=NAVY if b else '000000'); c.alignment = Alignment(wrap_text=True)

# summary
ws = wb.create_sheet('Summary')
hdr(ws, ['Grade', 'Unit', 'Domain', 'Weeks', 'Lessons S', 'Lessons P', 'Lessons I', 'Exam week', 'Report', 'Power standards', 'Test'],
    [14, 6, 34, 11, 9, 9, 9, 9, 10, 44, 50])
r = 2
for g in G.GRADES:
    P, units, order, dom_units, tests, mapwk, per = G.grade_model(g)
    for b, u in order:
        t = tests.get(u['dom']) if u['kind'] == 'domain' else None
        vals = [G.GNAME[g], 'D' + b if u['kind'] == 'domain' else 'E', u['n'], f"{u['a']}–{u['z']}",
                *[len([x for x in per[s] if x['unit'] == (b if u['kind'] == 'domain' else 'E')]) for s, _, _ in G.STRANDS],
                t['wk'] if t else '', ('Semester 1' if t['sem'] == 'S1' else 'Semester 2') if t else '', ', '.join(u['pw']), t['t'][0] if t and t.get('t') else '']
        for j, v in enumerate(vals, 1):
            c = ws.cell(r, j, v); c.font = Font(name=F, size=10); c.border = BOX; c.alignment = Alignment(wrap_text=True, vertical='top')
            c.fill = PatternFill('solid', fgColor=DCOL.get(u['dom'], 'FFFFFF'))
        if t and t.get('t'): ws.cell(r, 11).hyperlink = f"https://drive.google.com/file/d/{t['t'][1]}/view"
        r += 1
    c = ws.cell(r, 3, f'MAP Growth window {mapwk}'); c.font = Font(name=F, size=10, bold=True); c.fill = MAPF
    r += 2

# grade tabs
COLS = ['Week', 'Qtr', 'Week of', 'Days', 'Unit', 'Domain', 'Lesson', 'CCSS', 'Power', 'Type', 'From', 'S', 'P', 'I', 'Support blocks this week (pick one)']
for g in G.GRADES:
    P, units, order, dom_units, tests, mapwk, per = G.grade_model(g)
    ws = wb.create_sheet(G.GNAME[g])
    hdr(ws, COLS, [6, 5, 14, 5, 6, 30, 46, 22, 7, 22, 6, 4, 4, 4, 80])
    on = {s: {(x['wk'], x['n']) for x in per[s]} for s, _, _ in G.STRANDS}
    r = 2
    for w in P['weeks']:
        wk = w['wk']; first = True
        pl = '; '.join(f"{p.get('tag', '')} {p['name']}" for p in (w['t'].get('standard') or {}).get('pl') or [])
        stl = (w['t'].get('standard') or {}).get('st', [])
        extra = [x for s in ('priority', 'intervention') for x in per[s] if x['wk'] == wk and (wk, x['n']) not in on['standard']]
        seen = set()
        items = []
        for x in stl:
            items.append(x)
        for x in extra:
            if x['n'] not in seen: seen.add(x['n']); items.append(x)
        std_rows = {(x['wk'], x['n']): x for x in per['standard']}
        for x in items:
            k = G.lesson_kind(x)
            if k in ('exam', 'map'):
                t = tests.get(x.get('d')) if k == 'exam' else None
                vals = [wk, G.WD[wk]['q'], G.wdates(wk), G.WD[wk].get('days', 0), '', G.DOMSHORT.get(x.get('d'), '') if k == 'exam' else '',
                        ('DOMAIN EXAM · ' + G.DOMSHORT.get(x.get('d'), '')) if k == 'exam' else 'MAP GROWTH WINDOW (SPRING)', '', '', 'EXAM' if k == 'exam' else 'MAP', '', 'x', 'x', 'x', (t['t'][0] if t and t.get('t') else '')]
                fill = EXAMF if k == 'exam' else MAPF
            else:
                m = std_rows.get((wk, x['n'])) or next((y for s in ('priority', 'intervention') for y in per[s] if y['wk'] == wk and y['n'] == x['n']), x)
                ub = m.get('unit', '')
                u = units.get(ub, {})
                vals = [wk, G.WD[wk]['q'], G.wdates(wk), G.WD[wk].get('days', 0), ('D' + ub) if ub in dom_units else 'E', u.get('n', ''),
                        x['n'], ', '.join(x.get('c') or []), 'POWER' if x.get('p') else '', KIND[k], x.get('g', ''),
                        *['x' if (wk, x['n']) in on[s] else '' for s, _, _ in G.STRANDS], pl if first else '']
                fill = PatternFill('solid', fgColor=DCOL.get(u.get('dom', ''), 'FFFFFF'))
                first = False
            for j, v in enumerate(vals, 1):
                c = ws.cell(r, j, v); c.font = Font(name=F, size=10, bold=k in ('exam', 'map', 'build')); c.border = BOX
                c.alignment = Alignment(wrap_text=j in (6, 7, 8, 15), vertical='top'); c.fill = fill
            r += 1
    ws.auto_filter.ref = f'A1:{get_column_letter(len(COLS))}{ws.max_row}'

# coverage
ws = wb.create_sheet('CCSS coverage', 2)
hdr(ws, ['Grade', 'Standard', 'Power', 'First taught (Standard strand)', 'Before MAP?', 'Lessons that teach it (site)'], [14, 14, 8, 16, 10, 90])
r = 2
for g in G.GRADES:
    P, units, order, dom_units, tests, mapwk, per = G.grade_model(g)
    first = {}
    for x in per['standard']:
        if x['unit'] == 'E': continue
        for c in x.get('c') or []: first.setdefault(c, x['wk'])
    for d in G.CMAP['grades'][g]['domains']:
        for row in d['rows']:
            if not row['c'].startswith(G.PREFIX[g]): continue
            fw = first.get(row['c'], '')
            vals = [G.GNAME[g], row['c'], 'POWER' if row.get('pw') else '', fw or 'not on the Standard strand',
                    'yes' if fw and G.wn(fw) <= G.wn(mapwk) else 'NO', '; '.join(f"{l['t']} ({l['g']})" for l in row.get('ls', []))]
            for j, v in enumerate(vals, 1):
                c = ws.cell(r, j, v); c.font = Font(name=F, size=10, bold=(j == 3 and v == 'POWER')); c.border = BOX
                c.alignment = Alignment(wrap_text=j == 6, vertical='top'); c.fill = PatternFill('solid', fgColor=DCOL.get(d['d'], 'FFFFFF'))
            r += 1
ws.auto_filter.ref = f'A1:F{r - 1}'
out = sys.argv[2] if len(sys.argv) > 2 else 'out/docs/Awsaj-Domain-Sequence-K-5-2026-27.xlsx'
wb.save(out); print('saved', out, wb.sheetnames)
