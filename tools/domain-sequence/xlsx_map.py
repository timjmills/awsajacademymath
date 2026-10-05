import json
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment,Border,Side
from openpyxl.utils import get_column_letter
from core import GRADE_LABEL,GR,DOMNAME
A=json.load(open("final_map.json"))
F="Arial"; wb=Workbook()
NAVY="1F3864"
HDR=PatternFill("solid",fgColor=NAVY); HF=Font(name=F,bold=True,color="FFFFFF",size=10)
B=Font(name=F,size=10); BB=Font(name=F,size=10,bold=True)
TH=Side(style="thin",color="D0D0D0"); BOX=Border(left=TH,right=TH,top=TH,bottom=TH)
DCOL={"CC":"FFF2CC","OA":"DDEBF7","NBT":"E2EFDA","NF":"FCE4D6","MD":"EDE1F5","G":"DEF0F0","EXT":"EDEDED"}
MAPF=PatternFill("solid",fgColor="B4A7D6"); ENRF=PatternFill("solid",fgColor="F2F2F2")
EXAMF=PatternFill("solid",fgColor="FFD966")
SEMF={"S1":PatternFill("solid",fgColor="DDEBF7"),"S2":PatternFill("solid",fgColor="FCE4D6")}
def wrap(c,bold=False,size=10,color="000000"):
    c.font=Font(name=F,size=size,bold=bold,color=color); c.alignment=Alignment(wrap_text=True,vertical="top")

# ---------- READ ME
ws=wb.active; ws.title="READ ME"; ws.column_dimensions["A"].width=122
L=[("WRM reordered by CCSS domain, K to Grade 5 (MAP-ready, full CCSS coverage)",16,True,NAVY),("",10,False,""),
("Every White Rose small step for 2026-27, regrouped so each CCSS domain is taught as one run and examined once at the end of it, with every grade-level standard taught before the spring MAP Growth window. Built 5 October 2026 from WRM_to_CCSS_Mapping_v3 (828 tagged steps), the K-Grade 5 Weekly Pacing Guides (teaching order, week dates, CCSS BUILD lessons) and the Math Standards for Report Cards power-standards list.",10,False,""),("",10,False,""),
("HOW THE YEAR IS SHAPED",12,True,NAVY),
("1. Domain runs. Each grade-level domain is taught as one unbroken run and closed by a gold DOMAIN EXAM row.",10,False,""),
("2. MAP marker. A purple MAP GROWTH WINDOW row sits after the last domain exam. Everything above it is grade-level work.",10,False,""),
("3. Full CCSS coverage. All 148 K-5 CCSS standards are taught in their own grade before the end of April. Five standards were missing from their own grade's WRM guide (2.G.A.2, 3.OA.A.1, 3.OA.A.2, 4.MD.C.5, 5.MD.C.3); seven existing WRM steps from the neighbouring grade's guide are copied in to close them, placed at the right point in that domain's progression. They are marked COPIED IN. The CCSS coverage tab lists every standard and the week it is first taught.",10,False,""),
("4. Exam check. Every question in the six domain assessment booklets (305 questions) is tagged to its CCSS standard and checked against this sequence: 296 are fully taught before their domain exam, and every question's main grade-level standard is taught in time. The nine flagged questions, each with a fix, are on the Exam check tab.",10,False,""),
("5. Enrichment lessons. Every small step tagged only above grade level comes after the MAP marker, grouped by domain: WRM's year-ahead steps first, then Grade 6-only content (ratio, algebra, negative numbers). Steps that mix grade-level and above-grade codes stay in their domain run, because they teach grade-level content.",10,False,""),("",10,False,""),
("MAP READINESS: GRADE-LEVEL STANDARDS TAUGHT BY THE END OF APRIL (W32)",12,True,NAVY),
("MAP Growth reports four goal areas: Operations & Algebraic Thinking, Number & Operations (place value and fractions together), Measurement & Data, and Geometry.",10,False,"")]
for g in GR:
    cv=A[g]['cov']; b=A[g]['balance']['exam']
    txt="   ".join(f"{k} {v[0]}/{v[1]}" for k,v in cv.items())
    L.append((f"{GRADE_LABEL[g]:<28} {txt}      last domain exam W{A[g]['lastexam']:02d}   ·   power standards by exam S1 {b[0]} / S2 {b[1]}   ·   {A[g]['post']} enrichment steps",10,False,""))
L+=[("",10,False,""),("HOW THE ORDER WAS CHOSEN",12,True,NAVY),
("Each step goes to the domain of its grade-level CCSS code (the power-standard domain wins where a step carries two). Within a domain, steps keep White Rose's own order. Domains are ordered by prerequisite (counting before operations in K, place value before make-ten in Grades 1-2, operations before fractions, number before measurement); geometry has no prerequisites and is placed where it serves MAP and the S1/S2 balance. Among valid orders the plan picks, in priority: full MAP coverage, every domain exam before MAP, power standards balanced across S1 and S2, then the order closest to White Rose's.",10,False,""),("",10,False,""),
("WHAT TO CHECK BEFORE ADOPTING",12,True,NAVY),
("The spiral is gone. Once a domain is examined it is not taught again before MAP. Add a 5-10 minute daily retrieval slot that cycles all four MAP goal areas, not just number. MAP is adaptive, so students working below grade level will be served earlier-grade items from every goal area.",10,False,""),
("Grade 4 is the tightest grade: its last domain (Measurement & Data) is examined in W34, the week of 9 May. If MAP runs in the first week of May, Grade 4 should test at the end of the window.",10,False,""),
("WRM end-of-block assessments no longer match the teaching order. The slides, worksheets and Guide D teaching guides are per step and move with their step.",10,False,""),
("Why enrichment steps were not moved into the next grade: of the 165 above-grade steps that have a K-5 home, only seven were needed to close a gap. The rest repeat content the next grade's own guide already teaches, and moving them all would have pushed Grades 3 and 4 to about 5.4 steps a week before MAP. Nothing is missed by leaving them as enrichment: the next grade teaches that content in full. The 32 Grade 6-only steps (ratio, algebra, negative numbers) have no K-5 home and stay as enrichment.",10,False,""),
("Rows marked BUILD are CCSS lessons still to be written. COPIED IN steps use the slides and worksheet from the source grade's lesson folder.",10,False,""),
("Exams in W15 (week of 13 Dec) fall in the week S1 grades are due on 17 Dec; sit those early in the week.",10,False,""),
("Prerequisite judgements were made at domain level from CCSS and the step titles, not from WRM's own small-step dependency notes. Treat each grade as a draft for its band coordinator to check.",10,False,""),("",10,False,""),
("Q1 30 Aug - 22 Oct  |  Q2 1 Nov - 17 Dec -> S1 report  |  Q3 4 Jan - 30 Mar  |  Q4 4 Apr - 22 Jun -> S2 report",10,False,"")]
for i,(t,sz,bo,col) in enumerate(L,1):
    c=ws.cell(i,1,t); wrap(c,bo,sz,col or "000000")

# ---------- SUMMARY
ws=wb.create_sheet("Summary")
hdr=["Grade","Run","Domain","Steps","Taught weeks","Exam week","Exam week of","Semester","Power standards examined"]
widths=[24,5,34,7,13,10,14,10,62]
r=1
for g in GR:
    c=ws.cell(r,1,GRADE_LABEL[g]); wrap(c,True,12,NAVY)
    b=A[g]['balance']['exam']
    c=ws.cell(r,4,f"Power standards by exam: S1 {b[0]} / S2 {b[1]}   ·   MAP-ready: all grade-level standards taught by W32   ·   MAP marker W{A[g]['mapweek']:02d}   ·   {A[g]['post']} enrichment steps after it"); wrap(c,True,10,"404040")
    ws.merge_cells(start_row=r,start_column=4,end_row=r,end_column=9)
    r+=1
    for j,h in enumerate(hdr,1):
        c=ws.cell(r,j,h); c.fill=HDR; c.font=HF; c.alignment=Alignment(wrap_text=True)
    r+=1
    for k,d in enumerate([d for d in A[g]['cal'] if d['dom']!="EXT"],1):
        vals=[GRADE_LABEL[g],k,d['name'],d['steps'],f"W{d['start']:02d}–W{d['end']:02d}",
              (f"W{d['exam']:02d}" if d['exam'] else "none (not examined)"),d['exam_date'],d['sem'],", ".join(d['power'])]
        for j,v in enumerate(vals,1):
            c=ws.cell(r,j,v); c.font=B; c.border=BOX; c.alignment=Alignment(wrap_text=True,vertical="top")
            c.fill=PatternFill("solid",fgColor=DCOL[d['dom']])
        if d['sem'] in SEMF: ws.cell(r,8).fill=SEMF[d['sem']]; ws.cell(r,8).font=BB
        r+=1
    r+=1
for j,w in enumerate(widths,1): ws.column_dimensions[get_column_letter(j)].width=w

# ---------- GRADE TABS
cols=["#","Week","Qtr","Week of","Sem","Dom","Domain","Small step","CCSS codes","Power standard","Type","Was (WRM)","Moved (wks)","Teaching note"]
cw=[5,6,5,12,5,6,26,46,22,16,24,9,8,60]
for g in GR:
    name="Kindergarten" if g=="K" else "Grade "+g[1]
    ws=wb.create_sheet(name)
    for j,h in enumerate(cols,1):
        c=ws.cell(1,j,h); c.fill=HDR; c.font=HF; c.alignment=Alignment(wrap_text=True,vertical="center")
    ws.row_dimensions[1].height=30
    prev_sem=None
    for i,row in enumerate(A[g]['rows'],2):
        is_exam=row[10]=="EXAM"; is_map=row[10]=="MAP"; is_enr=str(row[10]).startswith("ENRICHMENT")
        for j,v in enumerate(row,1):
            c=ws.cell(i,j,v); c.border=BOX; c.alignment=Alignment(wrap_text=(j in (7,8,9,11,14)),vertical="top")
            c.font=Font(name=F,size=10,bold=(is_exam or is_map),italic=is_enr,color=("595959" if is_enr else "000000"))
            c.fill=MAPF if is_map else EXAMF if is_exam else ENRF if is_enr else PatternFill("solid",fgColor=DCOL.get(row[5],"FFFFFF"))
        if not is_map: ws.cell(i,5).fill=SEMF[row[4]]
        if row[10].startswith("BUILD") or row[10].startswith("COPIED"): ws.cell(i,11).font=Font(name=F,size=10,bold=True,color="C55A11")
        if row[13]: ws.cell(i,14).font=Font(name=F,size=10,color="C00000" if ("later" in row[13] or "Needs" in row[13] or "Uses" in row[13] or "Graphing" in row[13]) else "404040")
        if prev_sem=="S1" and row[4]=="S2":
            for j in range(1,len(cols)+1):
                ws.cell(i,j).border=Border(left=TH,right=TH,bottom=TH,top=Side(style="thick",color="C00000"))
        prev_sem=row[4]
    for j,w in enumerate(cw,1): ws.column_dimensions[get_column_letter(j)].width=w
    ws.freeze_panes="A2"; ws.auto_filter.ref=f"A1:{get_column_letter(len(cols))}{ws.max_row}"
COV=json.load(open("ccss_cov.json"))
from powerstd import POWER
PW={c for g in POWER for c,_,_ in POWER[g]}
ws=wb.create_sheet("CCSS coverage",2)
hd=["Grade","CCSS standard","Power","First taught (week)","Week of","Semester","First small step that teaches it","Before MAP?"]
for j,h in enumerate(hd,1):
    c=ws.cell(1,j,h); c.fill=HDR; c.font=HF; c.alignment=Alignment(wrap_text=True)
r=2
for g in GR:
    wk={row[1]:(row[3],row[4]) for row in A[g]['rows']}
    for code,w,step in COV[g]:
        alias={"K.CC.B.5":"K.CC.B.6"}.get(code,code)
        vals=[GRADE_LABEL[g],code,"POWER" if (code in PW or alias in PW) else "",f"W{w:02d}" if w else "NOT TAUGHT",wk.get(w,("",""))[0],wk.get(w,("",""))[1],step,"yes" if w and w<=32 else "NO"]
        for j,v in enumerate(vals,1):
            c=ws.cell(r,j,v); c.font=Font(name=F,size=10,bold=(j==3 and v=="POWER")); c.border=BOX
            c.fill=PatternFill("solid",fgColor=DCOL.get(code.split('.')[1],"FFFFFF"))
        r+=1
for j,w in enumerate([24,14,8,12,13,9,52,10],1): ws.column_dimensions[get_column_letter(j)].width=w
ws.freeze_panes="A2"; ws.auto_filter.ref=f"A1:H{r-1}"
EC=json.load(open("examcheck_rows.json"))
ws=wb.create_sheet("Exam check",3)
hd=["Grade","Domain test","Q","What the question asks","Standard(s) assessed","First taught (week)","Domain exam","Status","Fix / note"]
for j,h in enumerate(hd,1):
    c=ws.cell(1,j,h); c.fill=HDR; c.font=HF; c.alignment=Alignment(wrap_text=True)
for i,row in enumerate(EC,2):
    for j,v in enumerate(row,1):
        c=ws.cell(i,j,v); c.font=Font(name=F,size=10,bold=(j==8)); c.border=BOX; c.alignment=Alignment(wrap_text=(j in(2,4,9)),vertical="top")
    ws.cell(i,8).fill=PatternFill("solid",fgColor=("E2EFDA" if row[7]=="OK" else "FCE4D6"))
for j,w in enumerate([24,28,5,46,22,16,10,8,70],1): ws.column_dimensions[get_column_letter(j)].width=w
ws.freeze_panes="A2"; ws.auto_filter.ref=f"A1:I{len(EC)+1}"
out="out/WRM_Domain_Sequence_K-5.xlsx"
wb.save(out); print("saved",out,wb.sheetnames)
